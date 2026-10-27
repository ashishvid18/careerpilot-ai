from pathlib import Path
import io


def extract_resume_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            import pymupdf as fitz  # preferred name in newer versions
        except ImportError:
            try:
                import fitz  # fallback for older installs
            except ImportError:
                raise RuntimeError(
                    "PDF parsing requires PyMuPDF. "
                    "Install it with: pip install PyMuPDF"
                )
        with fitz.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(str(page.get_text("text")) for page in pdf)
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError(
                "DOCX parsing requires python-docx. "
                "Install it with: pip install python-docx"
            )
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        parts += [" | ".join(c.text.strip() for c in row.cells)
                  for table in doc.tables for row in table.rows]
        text = "\n".join(parts)
    else:
        raise ValueError("Please upload a PDF or DOCX resume.")
    if len(text.strip()) < 40:
        raise ValueError("I couldn't read enough text. Try a text-based PDF or DOCX.")
    return text.strip()
